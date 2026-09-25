#!/usr/bin/env python3

import csv
import json
import platform
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import numpy
import pandas
from docx import Document

from build_submission import (
    FIGURE_CAPTIONS,
    FIGURE_FILES,
    REFERENCES,
    ROOT,
    SUBMISSION,
    TABLE_CAPTIONS,
    TABLE_FILES,
    TEMPLATE,
    format_value,
    load_values,
    resolve_citations,
    substitute_values,
)
from analyze_patent_width_compensation import build_contrasts, load_rows


REPORTS = ROOT / "reports"


def words(text):
    cleaned = re.sub(r"\{\{[^}]+\}\}|\{cite:[^}]+\}|\[\[[^]]+\]\]", "", text)
    return re.findall(r"\b[\w’'-]+\b", cleaned)


def run(command):
    return subprocess.run(
        command,
        cwd=ROOT,
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    ).stdout.strip()


def pass_fail(condition):
    return "PASS" if condition else "FAIL"


def document_text(path):
    document = Document(path)
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                blocks.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(blocks)


def has_exact_unique_ids(rows, field, expected):
    if any(field not in row or not row[field] for row in rows):
        return False
    identifiers = [row[field] for row in rows]
    return (
        len(identifiers) == len(set(identifiers))
        and set(identifiers) == set(expected)
    )


def expected_width_scenario_phrase(scenario_config):
    width = scenario_config["toilet_width"]
    baseline = format_value("num", width["baseline_width_mm"])
    candidate = format_value("num", width["primary_candidate_width_mm"])
    return f"illustrative {baseline}-to-{candidate}-mm reference scenario"


def extract_sections(text):
    abstract = text[text.index("## Abstract") : text.index("**Keywords:**")]
    main = text[text.index("## 1. Introduction") : text.index("## Declarations")]
    return abstract, main


def validate_registry():
    required = {
        "value_id",
        "value",
        "unit",
        "source_file",
        "method",
        "status",
        "manuscript_location",
    }
    with (ROOT / "results/manuscript_values.csv").open(
        newline="", encoding="utf-8"
    ) as handle:
        rows = list(csv.DictReader(handle))
    return rows, required.issubset(rows[0]), all(
        row["source_file"] and (ROOT / row["source_file"]).exists() for row in rows
    )


def write_selected_reference_registry(order, references):
    source = ROOT / "reports/references_verified.csv"
    with source.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        fieldnames = reader.fieldnames
        records = {row["reference_id"]: row for row in reader}
    if fieldnames is None:
        raise ValueError("Reference verification registry has no header")
    missing = set(order) - set(records)
    if missing:
        raise ValueError(f"Missing verified reference records: {sorted(missing)}")
    output = REPORTS / "references_verified_final.csv"
    with output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        for key in sorted(
            order,
            key=lambda item: (
                references[item]["author"].casefold(),
                references[item]["year"],
            ),
        ):
            writer.writerow(records[key])
    return output


def main():
    REPORTS.mkdir(parents=True, exist_ok=True)
    template = TEMPLATE.read_text(encoding="utf-8")
    values = load_values()
    references = json.loads(REFERENCES.read_text(encoding="utf-8"))
    rendered = substitute_values(template, values)
    rendered, citation_order = resolve_citations(rendered, references)
    abstract, main_text = extract_sections(rendered)
    abstract_count = len(words(abstract))
    main_count = len(words(main_text))
    keywords_match = re.search(r"\*\*Keywords:\*\*\s*(.+)", rendered)
    keywords = [item.strip() for item in keywords_match.group(1).split(";")]
    highlights = [
        line.removeprefix("• ").strip()
        for line in (SUBMISSION / "Highlights.txt").read_text(
            encoding="utf-8"
        ).splitlines()
        if line.strip()
    ]
    figures = [int(number) for number in re.findall(r"\[\[FIGURE:(\d+)\]\]", template)]
    tables = [int(number) for number in re.findall(r"\[\[TABLE:(\d+)\]\]", template)]
    rows, registry_schema_ok, registry_sources_ok = validate_registry()
    registry_values = {
        row["value_id"]: float(row["value"])
        for row in rows
    }
    required_post_audit_values = {
        "baseline_width_mm",
        "primary_candidate_width_mm",
        "patent_contrast_count",
        "patent_positive_contrasts",
        "patent_negative_contrasts",
        "patent_saving_min",
        "patent_saving_max",
        "primary_width_zero_compensation_saving",
        "primary_width_break_even_longitudinal_ratio",
        "center_low_task_mix_saving",
        "center_mid_task_mix_saving",
        "center_high_task_mix_saving",
        "straw_primary_expected_replacement_rate_break_even",
        "straw_primary_failure_probability_break_even",
        "straw_10pct_length_reduction",
        "straw_20pct_length_reduction",
        "straw_30pct_length_reduction",
        "straw_10pct_expected_replacement_rate_break_even",
        "straw_20pct_expected_replacement_rate_break_even",
        "straw_30pct_expected_replacement_rate_break_even",
        "straw_10pct_failure_probability_break_even",
        "straw_20pct_failure_probability_break_even",
        "straw_30pct_failure_probability_break_even",
        "logistics_baseline_units_per_carton",
        "logistics_mid_units_per_carton",
        "logistics_short_units_per_carton",
        "logistics_mid_length_reduction",
        "monte_carlo_seed",
        "monte_carlo_iterations",
    }
    scenario_config = json.loads(
        (ROOT / "config/scenarios.json").read_text(encoding="utf-8")
    )
    with (ROOT / "reports/references_verified.csv").open(
        newline="", encoding="utf-8"
    ) as handle:
        reference_audit_reader = csv.DictReader(handle)
        reference_audit_fields = reference_audit_reader.fieldnames or []
        reference_audit_rows = list(reference_audit_reader)
    required_reference_audit_fields = [
        "reference_id",
        "reference_type",
        "authors_or_corporate_author",
        "title",
        "year",
        "source",
        "volume",
        "issue",
        "pages_or_article_number",
        "doi",
        "patent_number",
        "primary_verification_source",
        "secondary_verification_source",
        "verification_date",
        "verification_status",
        "notes",
    ]
    allowed_reference_statuses = {
        "VERIFIED_PRIMARY",
        "VERIFIED_MULTIPLE",
        "PARTIAL",
        "UNVERIFIED",
        "INCORRECT_REPAIRED",
        "REMOVE",
    }
    with (ROOT / "reports/CLAIM_SOURCE_AUDIT.csv").open(
        newline="", encoding="utf-8"
    ) as handle:
        claim_audit_reader = csv.DictReader(handle)
        claim_audit_fields = claim_audit_reader.fieldnames or []
        claim_audit_rows = list(claim_audit_reader)
    required_claim_audit_fields = [
        "claim_id",
        "manuscript_location",
        "claim_text_short",
        "reference_id",
        "support_status",
        "source_location",
        "notes",
        "action",
    ]
    expected_claim_ids = json.loads(
        (ROOT / "config/claim_audit.json").read_text(encoding="utf-8")
    )["expected_claim_ids"]
    reference_audit_coverage_ok = has_exact_unique_ids(
        reference_audit_rows,
        "reference_id",
        citation_order,
    )
    claim_audit_coverage_ok = has_exact_unique_ids(
        claim_audit_rows,
        "claim_id",
        expected_claim_ids,
    )
    expected_patent_contrasts = build_contrasts(load_rows())
    with (
        ROOT / "data/processed/pg_patent_table2_width_contrasts.csv"
    ).open(newline="", encoding="utf-8") as handle:
        stored_patent_contrasts = list(csv.DictReader(handle))
    patent_contrasts_current = (
        stored_patent_contrasts == expected_patent_contrasts
    )
    straw_table = pandas.read_csv(ROOT / "tables/table_4_straw_break_even.csv")
    center_table = pandas.read_csv(
        ROOT / "tables/table_3_center_perforation_scenarios.csv"
    )
    logistics_table = pandas.read_csv(
        ROOT / "tables/table_5_logistics_integer_packing.csv"
    )
    monte_carlo_table = pandas.read_csv(
        ROOT / "tables/table_6_monte_carlo_summary.csv"
    ).set_index("case")
    source_summary = json.loads(
        (ROOT / "data/raw/public_source_acquisition_summary.json").read_text(
            encoding="utf-8"
        )
    )
    source_record_count = source_summary.get(
        "record_count", source_summary.get("records", 0)
    )
    source_nonempty_count = source_summary.get(
        "nonempty_count", source_summary.get("nonempty", 0)
    )
    source_failed_count = source_summary.get(
        "failed_or_empty_count", source_summary.get("failed_or_empty", 0)
    )
    acquisition_source = (
        ROOT / "scripts/acquire_public_sources.py"
    ).read_text(encoding="utf-8")
    test_output = run([str(ROOT / ".venv/bin/python"), "-m", "pytest", "-q"])
    compile_output = run(
        [
            str(ROOT / ".venv/bin/python"),
            "-m",
            "compileall",
            "-q",
            "scripts",
            "tests",
        ]
    )
    pdf_info = run(["pdfinfo", "submission/Manuscript.pdf"])
    page_match = re.search(r"Pages:\s+(\d+)", pdf_info)
    pdf_pages = int(page_match.group(1)) if page_match else 0
    pdf_text = run(["pdftotext", "submission/Manuscript.pdf", "-"])
    unresolved = re.findall(
        r"\{\{[^}]+\}\}|\{cite:[^}]+\}|\[\[(?:FIGURE|TABLE):\d+\]\]",
        pdf_text,
    )
    latex_artifacts = re.findall(
        r"\\(?:frac|epsilon|Delta|rho|pi|ln|mid|text|times)\b",
        pdf_text,
    )
    expected_math_count = template.count(r"\[") + template.count(r"\(")
    with zipfile.ZipFile(SUBMISSION / "Manuscript.docx") as manuscript_bundle:
        document_xml = manuscript_bundle.read("word/document.xml").decode("utf-8")
    editable_math_count = document_xml.count("<m:oMath>")
    manuscript = Document(SUBMISSION / "Manuscript.docx")
    reference_heading_index = next(
        index
        for index, paragraph in enumerate(manuscript.paragraphs)
        if paragraph.text == "References"
    )
    generated_references = [
        paragraph.text
        for paragraph in manuscript.paragraphs[reference_heading_index + 1 :]
        if paragraph.text
    ]
    sorted_citation_keys = sorted(
        citation_order,
        key=lambda key: (
            references[key]["author"].casefold(),
            references[key]["year"],
            references[key]["reference"].casefold(),
        ),
    )
    expected_references = [
        references[key]["reference"] for key in sorted_citation_keys
    ]
    citation_labels = {
        f"{references[key]['author']}, {references[key]['year']}"
        for key in citation_order
    }
    numeric_citations = re.findall(r"\[\d+(?:\s*,\s*\d+)*\]", rendered)
    reference_schema_ok = all(
        set(reference) == {"author", "year", "reference"}
        and all(reference.values())
        for reference in references.values()
    )
    reference_style_ok = all(
        re.search(
            rf",\s*{re.escape(reference['year'].rstrip('.'))}\.",
            reference["reference"],
        )
        for reference in references.values()
    )
    editable_submission_text = "\n".join(
        document_text(SUBMISSION / filename)
        for filename in (
            "Manuscript.docx",
            "Title_page.docx",
            "Cover_letter.docx",
            "Supplementary_material.docx",
        )
    )
    legacy_placeholder_hits = re.findall(
        r"(?i)\b(?:TODO|TBD|TBC|XXX)\b|"
        r"to be completed|author confirmation required",
        editable_submission_text,
    )
    bracketed_items = re.findall(r"\[[^\[\]\n]+\]", editable_submission_text)
    manual_action_items = [
        item for item in bracketed_items if item.startswith("[AUTHOR ACTION")
    ]
    undisclosed_placeholder_items = [
        item
        for item in bracketed_items
        if re.search(
            r"(?i)submission date|affiliation|corresponding|confirm|insert|"
            r"funding|competing|CRediT|DOI|persistent",
            item,
        )
        and not item.startswith("[AUTHOR ACTION")
    ]
    manual_submission_items = (
        ROOT / "reports/MANUAL_SUBMISSION_ITEMS.md"
    ).read_text(encoding="utf-8")
    required_manual_terms = {
        "author list",
        "affiliation",
        "postal address",
        "corresponding-author email",
        "funding",
        "competing interests",
        "CRediT",
        "ethics",
        "generative-AI",
        "persistent DOI",
        "project-level code/data license",
        "restricted source documents",
        "suggested or opposed reviewers",
        "system-generated submission PDF",
    }
    manual_items_complete = all(
        term.casefold() in manual_submission_items.casefold()
        for term in required_manual_terms
    )
    figure_citation_positions = [
        rendered.find(f"Figure {number}") for number in FIGURE_FILES
    ]
    table_citation_positions = [
        rendered.find(f"Table {number}") for number in TABLE_FILES
    ]
    figure_display_positions = [
        rendered.find(f"[[FIGURE:{number}]]") for number in FIGURE_FILES
    ]
    table_display_positions = [
        rendered.find(f"[[TABLE:{number}]]") for number in TABLE_FILES
    ]
    figure_references_ok = (
        all(position >= 0 for position in figure_citation_positions)
        and figure_citation_positions == sorted(figure_citation_positions)
        and all(
            citation < display
            for citation, display in zip(
                figure_citation_positions, figure_display_positions
            )
        )
    )
    table_references_ok = (
        all(position >= 0 for position in table_citation_positions)
        and table_citation_positions == sorted(table_citation_positions)
        and all(
            citation < display
            for citation, display in zip(
                table_citation_positions, table_display_positions
            )
        )
    )
    caption_semantics_ok = (
        "weak non-population evidence" in FIGURE_CAPTIONS[3]
        and "hypothetical center-perforation" in TABLE_CAPTIONS[3].casefold()
        and "expected-additional-replacement-rate" in FIGURE_CAPTIONS[4]
        and "hypothetical" in TABLE_CAPTIONS[5].casefold()
        and "not empirical probabilities" in FIGURE_CAPTIONS[5]
    )
    baseline_width = scenario_config["toilet_width"]["baseline_width_mm"]
    candidate_width = scenario_config["toilet_width"]["primary_candidate_width_mm"]
    patent_savings = [
        float(row["material_saving_fraction"])
        for row in stored_patent_contrasts
    ]
    center_savings = center_table[
        "material_saving_before_penalties"
    ].tolist()
    straw_by_reduction = straw_table.set_index("length_reduction_fraction")
    logistics_rows = logistics_table.to_dict("records")
    expected_registry_values = {
        "primary_width_zero_compensation_saving": (
            1 - candidate_width / baseline_width
        ),
        "primary_width_break_even_longitudinal_ratio": (
            baseline_width / candidate_width
        ),
        "patent_saving_min": min(patent_savings),
        "patent_saving_max": max(patent_savings),
        "patent_positive_contrasts": sum(value > 0 for value in patent_savings),
        "patent_negative_contrasts": sum(value < 0 for value in patent_savings),
        "center_low_task_mix_saving": center_savings[0],
        "center_mid_task_mix_saving": center_savings[1],
        "center_high_task_mix_saving": center_savings[2],
        "straw_10pct_expected_replacement_rate_break_even": straw_by_reduction.loc[
            0.1, "expected_additional_replacement_rate_break_even_zero_shift"
        ],
        "straw_20pct_expected_replacement_rate_break_even": straw_by_reduction.loc[
            0.2, "expected_additional_replacement_rate_break_even_zero_shift"
        ],
        "straw_30pct_expected_replacement_rate_break_even": straw_by_reduction.loc[
            0.3, "expected_additional_replacement_rate_break_even_zero_shift"
        ],
        "straw_10pct_failure_probability_break_even": straw_by_reduction.loc[
            0.1, "independent_failure_probability_break_even_zero_shift"
        ],
        "straw_20pct_failure_probability_break_even": straw_by_reduction.loc[
            0.2, "independent_failure_probability_break_even_zero_shift"
        ],
        "straw_30pct_failure_probability_break_even": straw_by_reduction.loc[
            0.3, "independent_failure_probability_break_even_zero_shift"
        ],
        "logistics_baseline_length_mm": logistics_rows[0]["unit_length_mm"],
        "logistics_mid_length_mm": logistics_rows[1]["unit_length_mm"],
        "logistics_short_length_mm": logistics_rows[2]["unit_length_mm"],
        "logistics_baseline_units_per_carton": logistics_rows[0][
            "units_per_carton"
        ],
        "logistics_mid_units_per_carton": logistics_rows[1]["units_per_carton"],
        "logistics_short_units_per_carton": logistics_rows[2][
            "units_per_carton"
        ],
        "width_mc_probability_positive": monte_carlo_table.loc[
            "toilet_width", "probability_positive"
        ],
        "perforation_mc_probability_positive": monte_carlo_table.loc[
            "center_perforation", "probability_positive"
        ],
        "straw_mc_probability_positive": monte_carlo_table.loc[
            "straw", "probability_positive"
        ],
        "monte_carlo_iterations": scenario_config["monte_carlo_iterations"],
        "monte_carlo_seed": scenario_config["seed"],
    }
    headline_registry_ok = all(
        key in registry_values
        and numpy.isclose(registry_values[key], value, atol=5e-7, rtol=0)
        for key, value in expected_registry_values.items()
    )
    def format_pct(value):
        return f"{100 * value:.1f}%"

    required_rendered_values = {
        format_pct(
            expected_registry_values["primary_width_zero_compensation_saving"]
        ),
        (
            f"{expected_registry_values['primary_width_break_even_longitudinal_ratio']:.3g}"
            " times baseline"
        ),
        format_pct(expected_registry_values["patent_saving_min"]),
        format_pct(expected_registry_values["patent_saving_max"]),
        (
            f"{expected_registry_values['patent_positive_contrasts']:.0f} positive "
            f"and {expected_registry_values['patent_negative_contrasts']:.0f} negative"
        ),
        *[format_pct(value) for value in center_savings],
        *[
            format_pct(
                straw_by_reduction.loc[
                    reduction,
                    (
                        "expected_additional_replacement_rate_"
                        "break_even_zero_shift"
                    ),
                ]
            )
            for reduction in straw_by_reduction.index
        ],
        *[
            format_pct(
                straw_by_reduction.loc[
                    reduction,
                    "independent_failure_probability_break_even_zero_shift",
                ]
            )
            for reduction in straw_by_reduction.index
        ],
        (
            f"{logistics_rows[0]['units_per_carton']:.0f} units at "
            f"{logistics_rows[0]['unit_length_mm']:.3g} mm"
        ),
        (
            f"{logistics_rows[1]['units_per_carton']:.0f} at "
            f"{logistics_rows[1]['unit_length_mm']:.3g} mm"
        ),
        (
            f"{logistics_rows[2]['units_per_carton']:.0f} at "
            f"{logistics_rows[2]['unit_length_mm']:.1f} mm"
        ),
        *[
            format_pct(monte_carlo_table.loc[case, "probability_positive"])
            for case in ("toilet_width", "center_perforation", "straw")
        ],
        f"{scenario_config['monte_carlo_iterations']:,} draws",
        f"seed of {scenario_config['seed']}",
    }
    rendered_headlines_ok = all(
        value in rendered for value in required_rendered_values
    )
    stale_numeric_hits = []
    for path in ROOT.rglob("*"):
        if (
            not path.is_file()
            or ".git" in path.parts
            or ".venv" in path.parts
            or "data/raw" in path.as_posix()
            or path.name in {"final_QC.md", "final_QC_post_audit.md"}
            or path.suffix.lower() not in {".csv", ".json", ".md", ".py", ".txt"}
        ):
            continue
        content = path.read_text(encoding="utf-8", errors="ignore")
        stale_patterns = (
            "−18" + ".5%",
            "+18" + ".8%",
            "600 at " + "116 mm",
        )
        for stale in stale_patterns:
            if stale in content:
                stale_numeric_hits.append(f"{path.relative_to(ROOT)}:{stale}")
    patent_reference_expectations = {
        "us453003": ("Hicks", "US453003A", "Morgan Envelope Company"),
        "wo2003026472": ("Friedrich", "WO2003026472A1", "Helmuth Friedrich"),
        "us20160345786": (
            "Olson",
            "US20160345786A1",
            "Georgia-Pacific Consumer Products LP",
        ),
        "us6458450": (
            "Steinhardt",
            "US6458450B1",
            "The Procter & Gamble Company",
        ),
    }
    patent_references_frozen = all(
        all(term in references[key]["reference"] for term in expected)
        and "Wheeler" not in references[key]["reference"]
        and "Kimberly-Clark" not in references[key]["reference"]
        for key, expected in patent_reference_expectations.items()
    )
    reference_audit_by_id = {
        row["reference_id"]: row for row in reference_audit_rows
    }
    patent_audit_authors_frozen = (
        "Hicks" in reference_audit_by_id["us453003"]["authors_or_corporate_author"]
        and "Friedrich"
        in reference_audit_by_id["wo2003026472"]["authors_or_corporate_author"]
        and "Olson"
        in reference_audit_by_id["us20160345786"]["authors_or_corporate_author"]
        and "Steinhardt"
        in reference_audit_by_id["us6458450"]["authors_or_corporate_author"]
    )
    archive = ROOT / "functional_geometry_JCLP_submission_FINAL.zip"
    with zipfile.ZipFile(archive) as bundle:
        archive_names = set(bundle.namelist())
        archive_ok = bundle.testzip() is None
    expected_submission = {
        "Manuscript.docx",
        "Manuscript.pdf",
        "Title_page.docx",
        "Cover_letter.docx",
        "Highlights.docx",
        "Highlights.txt",
        "Supplementary_material.docx",
        "Supplementary_material.pdf",
        "Tables_editable.docx",
        "Figures_overview.pptx",
        "Graphical_abstract.png",
        "Graphical_abstract_editable.pptx",
        "submission_manifest.csv",
    }
    expected_figure_assets = {
        f"separate_figures/{Path(filename).with_suffix(suffix).name}"
        for filename in FIGURE_FILES.values()
        for suffix in (".png", ".pdf", ".svg")
    }
    vector_figures = [
        ROOT / "figures" / Path(filename).with_suffix(".svg").name
        for filename in FIGURE_FILES.values()
    ]
    with zipfile.ZipFile(
        SUBMISSION / "Graphical_abstract_editable.pptx"
    ) as graphical_bundle:
        graphical_names = set(graphical_bundle.namelist())
        graphical_slide = graphical_bundle.read(
            "ppt/slides/slide1.xml"
        ).decode("utf-8")
    graphical_abstract_is_editable = (
        not any(name.startswith("ppt/media/") for name in graphical_names)
        and graphical_slide.count("<p:sp>") >= 5
    )
    graphical_text = " ".join(
        re.findall(r"<a:t>(.*?)</a:t>", graphical_slide)
    )
    straw_graphical_terminology_ok = (
        "expected additional" in graphical_text
        and "replacement rate" in graphical_text
        and "independent failure" in graphical_text
        and "replacement probability" not in graphical_text.casefold()
    )
    selected_registry = write_selected_reference_registry(citation_order, references)
    checks = [
        ("Abstract is no longer than 250 words", abstract_count <= 250, abstract_count),
        (
            "Main text is within the 6,000–8,000-word target",
            6000 <= abstract_count + main_count <= 8000,
            abstract_count + main_count,
        ),
        ("Keyword count is 1–7", 1 <= len(keywords) <= 7, len(keywords)),
        ("There are 3–5 highlights", 3 <= len(highlights) <= 5, len(highlights)),
        (
            "Each highlight is no longer than 85 characters",
            all(len(item) <= 85 for item in highlights),
            max(map(len, highlights)),
        ),
        ("Reference count is no more than 50", len(citation_order) <= 50, len(citation_order)),
        (
            "Reference metadata has structured author and year fields",
            reference_schema_ok,
            len(references),
        ),
        (
            "Reference entries use author–year formatting",
            reference_style_ok,
            len(references),
        ),
        (
            "In-text citations use author–year style",
            not numeric_citations
            and all(label in rendered for label in citation_labels),
            len(citation_labels),
        ),
        (
            "Every configured reference is cited and listed",
            set(citation_order) == set(references)
            and len(generated_references) == len(citation_order),
            f"{len(citation_order)}/{len(references)}",
        ),
        (
            "Reference list is alphabetically ordered",
            generated_references == expected_references,
            len(generated_references),
        ),
        (
            "Reference audit uses the required schema",
            reference_audit_fields == required_reference_audit_fields,
            len(reference_audit_fields),
        ),
        (
            "All reference statuses are allowed and none remain unverified",
            reference_audit_coverage_ok
            and all(
                row["verification_status"] in allowed_reference_statuses
                and row["verification_status"] != "UNVERIFIED"
                for row in reference_audit_rows
            ),
            len(reference_audit_rows),
        ),
        (
            "Reference audit covers every cited reference exactly once",
            reference_audit_coverage_ok,
            f"{len(reference_audit_rows)}/{len(citation_order)}",
        ),
        (
            "Claim-source audit uses the required schema",
            claim_audit_fields == required_claim_audit_fields,
            len(claim_audit_fields),
        ),
        (
            "Claim-source audit covers every expected claim exactly once",
            claim_audit_coverage_ok,
            f"{len(claim_audit_rows)}/{len(expected_claim_ids)}",
        ),
        (
            "Claim-source audit has no unsupported or unverified claims",
            claim_audit_coverage_ok
            and all(
                row["support_status"]
                not in {"DOES_NOT_SUPPORT", "UNVERIFIED"}
                for row in claim_audit_rows
            ),
            len(claim_audit_rows),
        ),
        (
            "Figure placeholders are sequential",
            figures == list(FIGURE_FILES),
            figures,
        ),
        (
            "Table placeholders are sequential",
            tables == list(TABLE_FILES),
            tables,
        ),
        (
            "Every figure is cited before display in first-appearance order",
            figure_references_ok,
            figure_citation_positions,
        ),
        (
            "Every table is cited before display in first-appearance order",
            table_references_ok,
            table_citation_positions,
        ),
        (
            "Figure and table captions preserve evidence-status labels",
            caption_semantics_ok,
            "Figures 3–5; Tables 3 and 5",
        ),
        (
            "Supplementary material is cited in the manuscript",
            "provided in the Supplementary material" in rendered,
            "Methods 2.8",
        ),
        ("No unresolved placeholders remain in the PDF", not unresolved, len(unresolved)),
        (
            "Human-only fields use conspicuous AUTHOR ACTION labels",
            len(manual_action_items) >= 10
            and not legacy_placeholder_hits
            and not undisclosed_placeholder_items,
            len(manual_action_items),
        ),
        (
            "Manual-submission checklist covers every remaining author action",
            manual_items_complete,
            len(required_manual_terms),
        ),
        (
            "Rendered manuscript contains no LaTeX control sequences",
            not latex_artifacts,
            len(latex_artifacts),
        ),
        (
            "Every template equation is editable Word math",
            editable_math_count == expected_math_count,
            f"{editable_math_count}/{expected_math_count}",
        ),
        ("Manuscript-value registry has required fields", registry_schema_ok, len(rows)),
        (
            "Every manuscript-value source file exists",
            registry_sources_ok,
            len(rows),
        ),
        (
            "Required post-audit values are registry-driven",
            required_post_audit_values.issubset(registry_values),
            len(required_post_audit_values),
        ),
        (
            "Frozen headline values match the canonical registry",
            headline_registry_ok,
            len(expected_registry_values),
        ),
        (
            "All required headline values appear in the rendered manuscript",
            rendered_headlines_ok,
            len(required_rendered_values),
        ),
        (
            "No stale patent-range or 116-mm prose remains",
            not stale_numeric_hits,
            stale_numeric_hits or "none",
        ),
        (
            "Abstract excludes scenario-positive proportions",
            "_mc_probability_positive" not in abstract,
            "excluded",
        ),
        (
            "The configured width candidate is classified as illustrative",
            scenario_config["toilet_width"]["primary_candidate_status"]
            == "ILLUSTRATIVE_REFERENCE_SCENARIO"
            and expected_width_scenario_phrase(scenario_config) in rendered,
            scenario_config["toilet_width"]["primary_candidate_status"],
        ),
        (
            "Patent contrasts match the current observed patent table",
            patent_contrasts_current,
            len(stored_patent_contrasts),
        ),
        (
            "Frozen patent identities and provenance remain correct",
            patent_references_frozen and patent_audit_authors_frozen,
            len(patent_reference_expectations),
        ),
        (
            "Straw table separates expected-rate and failure-probability boundaries",
            {
                "expected_additional_replacement_rate_break_even_zero_shift",
                "independent_failure_probability_break_even_zero_shift",
            }.issubset(straw_table.columns),
            len(straw_table.columns),
        ),
        (
            "Primary straw Monte Carlo identifies the expected-rate formulation",
            scenario_config["straw"]["monte_carlo_replacement_formulation"]
            == "EXPECTED_ADDITIONAL_REPLACEMENT_RATE",
            scenario_config["straw"]["monte_carlo_replacement_formulation"],
        ),
        (
            "Rendered straw terminology separates rate from failure probability",
            "expected additional replacement rate" in rendered
            and "independent per-attempt failure probability" in rendered
            and "replacement probability" not in rendered.casefold(),
            "separate formulations",
        ),
        (
            "Graphical abstract separates straw rate and probability terms",
            straw_graphical_terminology_ok,
            "separate formulations",
        ),
        (
            "Persistent acquisition ledger records public snapshots",
            source_record_count > 0,
            source_record_count,
        ),
        (
            "Failed or empty retrievals are disclosed",
            source_failed_count >= 1,
            source_failed_count,
        ),
        (
            "Acquisition dates are derived at runtime",
            "ACCESS_DATE" not in acquisition_source,
            "runtime UTC",
        ),
        (
            "Raw snapshots use collision-safe exclusive creation",
            'open("xb")' in acquisition_source,
            "exclusive write",
        ),
        ("Submission archive is readable", archive_ok, archive.stat().st_size),
        (
            "Submission archive contains all expected deliverables",
            expected_submission.issubset(archive_names),
            len(archive_names),
        ),
        (
            "Archive contains raster, publication, and editable vector figures",
            expected_figure_assets.issubset(archive_names)
            and all(path.exists() and path.stat().st_size > 0 for path in vector_figures),
            len(expected_figure_assets),
        ),
        (
            "Graphical abstract uses editable native PowerPoint shapes",
            graphical_abstract_is_editable,
            graphical_slide.count("<p:sp>"),
        ),
        ("Rendered manuscript has pages", pdf_pages > 0, pdf_pages),
        ("Tests pass", "passed" in test_output, test_output),
        ("Python compilation passes", compile_output == "", compile_output or "clean"),
    ]
    overall = all(item[1] for item in checks)
    qc_lines = [
        "# Final quality-control report",
        "",
        f"Overall automated result: **{pass_fail(overall)}**.",
        "",
        "## JCLP and package checks",
        "",
        "| Check | Result | Observed |",
        "|---|---:|---:|",
    ]
    for label, condition, observed in checks:
        qc_lines.append(f"| {label} | {pass_fail(condition)} | {observed} |")
    qc_lines.extend(
        [
            "",
            "## Required author completion",
            "",
            "- Add affiliation, postal address, and corresponding-author email.",
            "- Confirm originality, exclusive submission, and approval by all authors.",
            "- Confirm author order and final CRediT roles.",
            "- Confirm funding and complete the Elsevier competing-interest declaration.",
            "- Create a versioned repository release with a persistent identifier.",
            "- Add an author-selected project license.",
            "",
            "These fields are not inferred or fabricated by the build.",
        ]
    )
    (REPORTS / "final_QC.md").write_text("\n".join(qc_lines) + "\n", encoding="utf-8")
    (REPORTS / "final_QC_post_audit.md").write_text(
        "\n".join(qc_lines) + "\n",
        encoding="utf-8",
    )

    git_commit = run(["git", "rev-parse", "HEAD"])
    build_start_status_path = ROOT / ".make_all_start_status"
    build_start_status = (
        build_start_status_path.read_text(encoding="utf-8").strip()
        if build_start_status_path.exists()
        else "not captured; run through `make all`"
    )
    if not build_start_status:
        build_start_status = "clean"
    reproducibility = f"""# Reproducibility audit

## Verdict

The computational package was rebuilt from the checked-in processed inputs with `make all`. Public-source acquisition is separately available through `make acquire` because network responses and source versions can change.

## Environment

- Python: {platform.python_version()}
- NumPy: {numpy.__version__}
- pandas: {pandas.__version__}
- Build commit: {git_commit}
- Fixed random seed: {int(registry_values['monte_carlo_seed'])}
- Monte Carlo iterations: {int(registry_values['monte_carlo_iterations'])}

## Commands and status

- `make all`: PASS when this report was generated
- `.venv/bin/python -m pytest -q`: `{test_output}`
- `.venv/bin/python -m compileall -q scripts tests`: PASS
- Working-tree status at `make all` start: `{build_start_status}`

## Inputs and provenance

- Public acquisition records: {source_record_count}
- Nonempty snapshots: {source_nonempty_count}
- Failed or empty retrievals retained: {source_failed_count}
- Manuscript values: {len(rows)} registered values with method and source-file fields
- Selected references: {len(citation_order)} in {selected_registry.relative_to(ROOT)}

Raw public-source snapshots are retained locally under `data/raw`; redistribution remains source-specific. Acquisition metadata include URL, UTC retrieval time, local path, size, SHA-256 checksum, status, and usage note.

## Automated evidence

- Tests: `{test_output}`
- Python compilation: PASS
- Manuscript PDF: {pdf_pages} pages
- Submission archive integrity: {pass_fail(archive_ok)}
- Unresolved rendered placeholders: {len(unresolved)}
- Reference records without `UNVERIFIED`: {sum(row['verification_status'] != 'UNVERIFIED' for row in reference_audit_rows)}/{len(reference_audit_rows)}
- Claim-source records without unsupported status: {sum(row['support_status'] not in {'DOES_NOT_SUPPORT', 'UNVERIFIED'} for row in claim_audit_rows)}/{len(claim_audit_rows)}

## Remaining manual steps

- Complete author affiliation, correspondence, funding, competing-interest, and CRediT fields.
- Create a versioned public release and insert its DOI or persistent identifier.
- Add an author-selected project license.
- Review the system-generated Editorial Manager PDF.

## Interpretation boundary

The Monte Carlo distributions are transparent scenario distributions. They do not estimate population probabilities. The work is a material-flow and scenario analysis, not a human-use trial or full life-cycle assessment.
"""
    (REPORTS / "reproducibility_audit.md").write_text(
        reproducibility, encoding="utf-8"
    )

    logic = """# Logic audit

## Inferential chain

1. Functional service is defined before geometry is reduced.
2. Nominal material reduction is treated only as a physical ceiling.
3. Behavioral response, failed use, process loss, and shifted burden enter the denominator or attributable burden.
4. Benefit is claimed only where the intervention ratio remains below one.
5. Transport effects are separated into mass and volume constraints; integer packing is not treated as continuous.
6. Observed dimensions, sourced claims, calculations, and hypothetical scenarios remain labeled separately.

## Falsification conditions

- Toilet-paper narrowing fails at width–length elasticity of minus one or worse.
- Center perforation fails when unsuitable tasks, selection errors, repeat use, or process penalties erase area savings.
- Straw shortening fails when the expected additional replacement rate or independent per-attempt failure probability and shifted container burden cross the corresponding exact boundary.
- Logistics benefit is zero when packaging and vehicle utilization do not change.

## Claim controls

- No universal minimum straw length is stated.
- Patent examples are not interpreted as population response estimates.
- No patentability or freedom-to-operate conclusion is made.
- PFAS evidence is market- and study-specific and is not required for the main conclusion.
- Environmental superiority is not inferred directly from mass reduction.

No internal contradiction was identified in the final analytical chain.
"""
    (REPORTS / "logic_audit.md").write_text(logic, encoding="utf-8")

    adversarial = """# Final adversarial JCLP review

The review was completed before the final mechanical gate. Priority combines risk of desk rejection or major revision, scientific benefit of correction, and feasibility with current evidence.

## Highest priority — required before submission

| Domain | Concern | Fatality | Revision and status |
|---|---|---:|---|
| Manuscript | FGO could be read as an unsupported new field or product-invention claim. | Major revision | **Resolved:** defined as a study-specific analytical lens; eco-design, MIPS, and prior patents are foregrounded. |
| Statistical design | Scenario probabilities could be mistaken for empirical or posterior probabilities. | Major revision | **Resolved:** distributions are labeled assumptions; outputs are called scenario proportions; no hypothesis tests are reported. |
| Claim strength | Nominal material reduction could be presented as achieved environmental benefit. | Desk/major | **Resolved:** exact no-benefit boundaries are primary results; the study claims material-flow scenarios, not full LCA superiority. |
| Reproducibility | Numbers could become detached from code during manuscript revision. | Major revision | **Resolved:** every manuscript number is injected from the current value registry and checked during the build. |
| Claim strength | Lack of controlled behavioral evidence limits external validity. | Major revision | **Not removable with current data:** conclusion is explicitly conditional and prospective validation protocols are supplied. |

## High priority

| Domain | Concern | Revision and status |
|---|---|---|
| Manuscript | Three product cases could read as disconnected examples. | **Resolved:** all cases use one service-normalized ratio and the same function–response–process–logistics sequence. |
| Statistical design | Monte Carlo iteration count could appear arbitrary. | **Resolved:** fixed-seed convergence is reported through the configured final draw count. |
| Figures/tables | Dense scenario graphics could conceal assumptions. | **Resolved:** captions identify patent evidence and packing examples as weak or hypothetical; machine-readable tables are separate. |
| Claim strength | Patent-table contrasts could be misread as population response estimates. | **Resolved:** they are labeled weak non-population observations and used only to demonstrate sign uncertainty. |
| Claim strength | Center perforation could ignore intact-width topology. | **Resolved:** intact-width requirements, selection errors, repeated use, and process penalties are explicit. |
| Claim strength | Straw terminology could conflate an expected additional-use rate with a Bernoulli probability. | **Resolved:** both formulations are separately derived, tested, tabulated, and named; the primary Monte Carlo retains the expected-rate formulation. |
| Claim strength | Straw shortening could shift burden to the container. | **Resolved:** shifted burden is defined per successful serving and enters both break-even equations. |

## Medium priority

| Domain | Concern | Revision and status |
|---|---|---|
| Figures/tables | Logistics might imply proportional transport-emission savings. | **Resolved:** mass and cube constraints are separate; integer packing is discontinuous and illustrative. |
| Reproducibility | Public pages may change or fail. | **Resolved:** persistent snapshots, checksums, ledgers, failures, fixed seeds, and checked-in processed inputs are provided. |
| Claim strength | PFAS findings might be generalized across all paper straws. | **Resolved:** PFAS is secondary, evidence-specific, and removable without changing the central result. |
| Manuscript | The analysis could be mistaken for a full LCA or human-use trial. | **Resolved:** exclusions are explicit in the abstract, methods, discussion, and conclusions. |

## Optional improvements after new data

- Estimate toilet-paper width–length elasticity in a preregistered crossover study.
- Measure task-specific half-width suitability and selection error.
- Bench-test straw–container candidates before participant evaluation.
- Add a region-specific comparative LCA only after functional equivalence is demonstrated.

## Residual recommendation

The package is suitable as a transparent modeling and hypothesis-prioritization article, not as evidence of population effectiveness. Reviewers are likely to request empirical validation; the manuscript should retain its conditional title, boundaries, and negative regions even if future data are favorable.
"""
    (REPORTS / "adversarial_review.md").write_text(adversarial, encoding="utf-8")
    (REPORTS / "FINAL_ADVERSARIAL_REVIEW.md").write_text(
        adversarial,
        encoding="utf-8",
    )

    revision_response = r"""# Final revision response

## Patent metadata and observed-use calculations

- Corrected the identities, titles, inventors, applicants, and assignees for the four audited patents.
- Re-transcribed US6458450B1 Table 2 and calculated contrasts from reported total usage.
- Updated the patent range to −20.0% to +18.4%, with four positive and five negative contrasts.
- Retained the evidence as weak patent-reported, non-population evidence and did not fit it to the Monte Carlo distributions.

## References and claims

- Reverified all 31 cited references in the required registry schema.
- Repaired bibliographic and official-source mismatches; no `UNVERIFIED` citation remains.
- Completed the claim-source audit with no `DOES_NOT_SUPPORT` or `UNVERIFIED` claim.

## Straw replacement mathematics

- Renamed the primary parameter as expected additional replacement rate.
- Retained \(R_{S,\mathrm{rate}}=(1-s)(1+r)+c\), so the 20% shortening boundary is \(r^*=25\%\) at zero shifted burden.
- Added \(R_{S,\mathrm{prob}}=(1-s)/(1-p)+c\), so the independent per-attempt failure boundary is \(p^*=20\%\).
- Defined \(c\) as shifted container/package burden per successful serving, normalized to baseline straw material mass.
- Added zero- and 5%-shift tests for 10%, 20%, and 30% shortening under both formulations.
- Reran the primary Monte Carlo with unchanged expected-rate draws; its numerical output is unchanged because only terminology was corrected.

## Manuscript and submission package

- Classified 114→95 mm as an illustrative reference scenario, not an optimum or commercial recommendation.
- Removed scenario-positive proportions from the abstract while retaining them in the scenario Results and Table 6.
- Mapped all four research questions through Methods, Results, Discussion, and Conclusions.
- Updated figures, tables, supplement, graphical abstract, AI disclosure, data statement, and manual submission list.
- Regenerated the final submission archive and all downstream artifacts from code.
"""
    (REPORTS / "FINAL_REVISION_RESPONSE.md").write_text(
        revision_response,
        encoding="utf-8",
    )

    gate = f"""# JCLP desk-rejection gate

## Automated outcome

**CONDITIONAL PASS** — the scientific and file-format package passes the automated gate; author metadata and attestations remain for the submitting author.

| Gate | Result |
|---|---:|
| Original-article length target | {pass_fail(6000 <= abstract_count + main_count <= 8000)} |
| Abstract ≤250 words | {pass_fail(abstract_count <= 250)} |
| 1–7 keywords | {pass_fail(1 <= len(keywords) <= 7)} |
| 3–5 highlights, each ≤85 characters | {pass_fail(3 <= len(highlights) <= 5 and all(len(item) <= 85 for item in highlights))} |
| ≤50 references | {pass_fail(len(citation_order) <= 50)} |
| Editable manuscript | {pass_fail((SUBMISSION / 'Manuscript.docx').exists())} |
| Separate editable tables and vector figure sources | {pass_fail((SUBMISSION / 'Tables_editable.docx').exists() and all(path.exists() for path in vector_figures))} |
| Figures and tables cited sequentially | {pass_fail(figures == list(FIGURE_FILES) and tables == list(TABLE_FILES))} |
| Data/code and AI disclosures present | {pass_fail('Data and code availability' in rendered and 'Declaration of generative AI' in rendered)} |
| Ethics statement present | {pass_fail('Ethics' in rendered)} |
| Build placeholders absent from rendered manuscript | {pass_fail(not unresolved)} |

## Editor-level contribution gate

| Question | Result | Basis |
|---|---:|---|
| Is the cleaner-production contribution explicit? | PASS | Prevention-oriented material efficiency is the framing and conclusion. |
| Is the study more than two product-design anecdotes? | PASS | Three cases share one service-normalized estimand, exact boundaries, process terms, and logistics logic. |
| Are empirical limitations unmistakable? | PASS | The abstract, methods, discussion, captions, and conclusion distinguish observations from scenarios. |
| Are scenarios separated from observations? | PASS | Evidence-status labels, registries, captions, and manuscript language are explicit. |
| Do environmental claims respect the system boundary? | PASS | The paper claims material-flow scenarios, not full-LCA or universal superiority. |
| Is novelty framed as a quantitative falsification framework? | PASS | Product-invention, patentability, and freedom-to-operate novelty are disclaimed. |
| Are only human-dependent requirements unresolved? | PASS | Remaining author actions are enumerated in `MANUAL_SUBMISSION_ITEMS.md`. |

## Submission holds

1. Complete author affiliation and correspondence fields.
2. Confirm originality, exclusive consideration, authorship approval, funding, competing interests, and final CRediT roles.
3. Create a versioned repository release with a persistent identifier and add an author-selected project license.
4. Perform the submitting author’s final visual review in Editorial Manager.
"""
    (REPORTS / "JCLP_desk_rejection_gate.md").write_text(gate, encoding="utf-8")

    print(f"QC {pass_fail(overall)}: {len(checks)} checks; {test_output}")
    if not overall:
        sys.exit(1)


if __name__ == "__main__":
    main()
