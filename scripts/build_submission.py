#!/usr/bin/env python3

import csv
import json
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptInches
from pptx.util import Pt as PptPt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "manuscript/manuscript_template.md"
VALUES = ROOT / "results/manuscript_values.csv"
REFERENCES = ROOT / "config/references.json"
SUBMISSION = ROOT / "submission"

FIGURE_CAPTIONS = {
    1: (
        "Functional geometry optimization framework. Geometry is evaluated "
        "against functional constraints before response, process loss, and "
        "shifted burden determine material per successful service."
    ),
    2: (
        "Toilet-paper material-saving surface as a function of candidate width "
        "and width-compensation elasticity. The black boundary marks zero saving."
    ),
    3: (
        "Patent-table width contrasts and hypothetical center-perforation task "
        "mixtures. Patent contrasts are weak non-population evidence."
    ),
    4: (
        "Straw shortening expected-additional-replacement-rate boundary and an "
        "illustrative integer-packing example. Packing dimensions are hypothetical."
    ),
    5: (
        "Configured scenario uncertainty distributions. Densities and positive "
        "proportions are not empirical probabilities."
    ),
}

TABLE_CAPTIONS = {
    1: "Evidence domains, source counts, and principal limitations.",
    2: "Toilet-paper fixed-width scenarios and exact break-even ratios.",
    3: "Hypothetical center-perforation task-mixture scenarios.",
    4: "Straw-shortening break-even boundaries under two replacement formulations.",
    5: "Hypothetical integer-packing example.",
    6: "Monte Carlo scenario summaries.",
}

TABLE_FILES = {
    1: "table_1_evidence_base.csv",
    2: "table_2_width_scenarios.csv",
    3: "table_3_center_perforation_scenarios.csv",
    4: "table_4_straw_break_even.csv",
    5: "table_5_logistics_integer_packing.csv",
    6: "table_6_monte_carlo_summary.csv",
}

FIGURE_FILES = {
    1: "figure_1_fgo_framework.png",
    2: "figure_2_width_compensation_surface.png",
    3: "figure_3_observed_contrasts_and_perforation.png",
    4: "figure_4_straw_and_logistics_boundaries.png",
    5: "figure_5_scenario_uncertainty.png",
}


def load_values():
    with VALUES.open(newline="", encoding="utf-8") as handle:
        return {
            row["value_id"]: float(row["value"])
            for row in csv.DictReader(handle)
        }


def format_value(kind, value):
    if kind == "pct":
        return f"{100 * value:.1f}%"
    if kind == "int":
        return str(int(round(value)))
    if kind == "count":
        return f"{int(round(value)):,}"
    if kind == "num":
        return f"{value:.3g}"
    if kind == "num1":
        return f"{value:.1f}"
    raise ValueError(f"Unknown placeholder type: {kind}")


def substitute_values(text, values):
    pattern = re.compile(r"\{\{(pct|int|count|num|num1):([a-zA-Z0-9_]+)\}\}")

    def replacement(match):
        kind, key = match.groups()
        if key not in values:
            raise KeyError(f"Missing manuscript value: {key}")
        return format_value(kind, values[key])

    return pattern.sub(replacement, text)


def resolve_citations(text, reference_map):
    order = []

    def replacement(match):
        keys = [key.strip() for key in match.group(1).split(",")]
        for key in keys:
            if key not in reference_map:
                raise KeyError(f"Unknown citation key: {key}")
            if key not in order:
                order.append(key)
        labels = [
            f"{reference_map[key]['author']}, {reference_map[key]['year']}"
            for key in sorted(
                keys,
                key=lambda item: (
                    reference_map[item]["author"].casefold(),
                    reference_map[item]["year"],
                ),
            )
        ]
        return "(" + "; ".join(labels) + ")"

    return re.sub(r"\{cite:([^}]+)\}", replacement, text), order


def extract_braced(text, start):
    if start >= len(text) or text[start] != "{":
        raise ValueError(f"Expected braced expression in: {text}")
    depth = 0
    for index in range(start, len(text)):
        if text[index] == "{":
            depth += 1
        elif text[index] == "}":
            depth -= 1
            if depth == 0:
                return text[start + 1 : index], index + 1
    raise ValueError(f"Unclosed braced expression in: {text}")


def replace_latex_command(text, command, argument_count, formatter):
    marker = f"\\{command}"
    while marker in text:
        start = text.index(marker)
        cursor = start + len(marker)
        arguments = []
        for _ in range(argument_count):
            while cursor < len(text) and text[cursor].isspace():
                cursor += 1
            argument, cursor = extract_braced(text, cursor)
            arguments.append(argument)
        text = text[:start] + formatter(*arguments) + text[cursor:]
    return text


def latex_to_linear_math(latex):
    text = latex.strip()
    text = replace_latex_command(
        text,
        "frac",
        2,
        lambda numerator, denominator: (
            f"({latex_to_linear_math(numerator)})/"
            f"({latex_to_linear_math(denominator)})"
        ),
    )
    for command in ("text", "mathrm"):
        text = replace_latex_command(
            text,
            command,
            1,
            lambda value: latex_to_linear_math(value),
        )
    replacements = {
        r"\Delta": "Δ",
        r"\epsilon": "ε",
        r"\rho": "ρ",
        r"\pi": "π",
        r"\ln": "ln",
        r"\mid": "∣",
        r"\times": "×",
        r"\min": "min",
        r"\,": " ",
        r"\ ": " ",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    previous = None
    while previous != text:
        previous = text
        text = re.sub(r"_\{([^{}]+)\}", r"_(\1)", text)
        text = re.sub(r"\^\{([^{}]+)\}", r"^(\1)", text)
    text = text.replace("{", "(").replace("}", ")")
    unresolved = re.findall(r"\\[A-Za-z]+", text)
    if unresolved:
        raise ValueError(f"Unsupported LaTeX commands: {unresolved}")
    return re.sub(r"\s+", " ", text).strip()


def math_run(text):
    math_run = OxmlElement("m:r")
    run_properties = OxmlElement("m:rPr")
    style = OxmlElement("m:sty")
    style.set(qn("m:val"), "p")
    run_properties.append(style)
    math_run.append(run_properties)
    math_text = OxmlElement("m:t")
    math_text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    math_text.text = text
    math_run.append(math_text)
    return math_run


def scripted_math(base_nodes, subscript_nodes, superscript_nodes):
    if subscript_nodes and superscript_nodes:
        scripted = OxmlElement("m:sSubSup")
    elif subscript_nodes:
        scripted = OxmlElement("m:sSub")
    else:
        scripted = OxmlElement("m:sSup")
    base = OxmlElement("m:e")
    for node in base_nodes:
        base.append(node)
    scripted.append(base)
    if subscript_nodes:
        subscript = OxmlElement("m:sub")
        for node in subscript_nodes:
            subscript.append(node)
        scripted.append(subscript)
    if superscript_nodes:
        superscript = OxmlElement("m:sup")
        for node in superscript_nodes:
            superscript.append(node)
        scripted.append(superscript)
    return scripted


def parse_math_argument(text, cursor):
    if cursor >= len(text):
        raise ValueError(f"Missing script argument in: {text}")
    if text[cursor] == "(":
        nodes, cursor = parse_math_nodes(text, cursor + 1, ")")
        return nodes, cursor
    return [math_run(text[cursor])], cursor + 1


def parse_math_nodes(text, cursor=0, stop=None):
    nodes = []
    while cursor < len(text):
        if stop and text[cursor] == stop:
            return nodes, cursor + 1
        if text[cursor] == "(":
            inner, cursor = parse_math_nodes(text, cursor + 1, ")")
            base_nodes = [math_run("("), *inner, math_run(")")]
        else:
            base_nodes = [math_run(text[cursor])]
            cursor += 1
        subscript_nodes = []
        superscript_nodes = []
        while cursor < len(text) and text[cursor] in "_^":
            marker = text[cursor]
            argument, cursor = parse_math_argument(text, cursor + 1)
            if marker == "_":
                subscript_nodes = argument
            else:
                superscript_nodes = argument
        if subscript_nodes or superscript_nodes:
            nodes.append(
                scripted_math(base_nodes, subscript_nodes, superscript_nodes)
            )
        else:
            nodes.extend(base_nodes)
    if stop:
        raise ValueError(f"Unclosed math group in: {text}")
    return nodes, cursor


def append_omml(paragraph, latex):
    math = OxmlElement("m:oMath")
    linear = latex_to_linear_math(latex)
    nodes, _ = parse_math_nodes(linear)
    for node in nodes:
        math.append(node)
    paragraph._p.append(math)


def add_rich_text(paragraph, text):
    parts = re.split(r"(\\\(.*?\\\)|\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith(r"\(") and part.endswith(r"\)"):
            append_omml(paragraph, part[2:-2])
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    properties.append(element)


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size, color in [
        ("Title", 18, "17324D"),
        ("Heading 1", 14, "17324D"),
        ("Heading 2", 12, "1F5A7A"),
        ("Heading 3", 11, "2F7D57"),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True


def add_title_page(document):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(90)
    run = paragraph.add_run(
        "Redesign before substitution: functional geometry optimization "
        "of disposable products"
    )
    run.bold = True
    run.font.size = Pt(19)
    run.font.color.rgb = RGBColor(23, 50, 77)
    author = document.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_before = Pt(28)
    author.add_run("Tatsuki Onishi").bold = True
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.add_run(
        "[AUTHOR ACTION BEFORE SUBMISSION: add affiliation, postal address, "
        "and corresponding-author email.]"
    ).italic = True
    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_before = Pt(24)
    metadata.add_run(
        "Article type: Original article\n"
        "Target journal: Journal of Cleaner Production\n"
        "Word count and compliance details are provided in reports/final_QC.md"
    )
    document.add_page_break()


def add_table(document, number):
    caption = document.add_paragraph()
    caption.paragraph_format.space_before = Pt(14)
    caption.paragraph_format.space_after = Pt(5)
    run = caption.add_run(f"Table {number}. {TABLE_CAPTIONS[number]}")
    run.bold = True
    frame = pd.read_csv(ROOT / "tables" / TABLE_FILES[number])
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, column in enumerate(frame.columns):
        header.cells[index].text = column.replace("_", " ")
        set_cell_shading(header.cells[index], "D9EAF2")
        for paragraph in header.cells[index].paragraphs:
            for cell_run in paragraph.runs:
                cell_run.bold = True
                cell_run.font.size = Pt(8)
    for row in frame.itertuples(index=False):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if isinstance(value, float):
                text = f"{value:.4g}"
            else:
                text = str(value)
            cells[index].text = text
            for paragraph in cells[index].paragraphs:
                for cell_run in paragraph.runs:
                    cell_run.font.size = Pt(7.5)


def add_figure(document, number):
    image = ROOT / "figures" / FIGURE_FILES[number]
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.add_run().add_picture(str(image), width=Inches(6.5))
    caption = document.add_paragraph()
    caption.paragraph_format.space_before = Pt(12)
    caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = caption.add_run(f"Figure {number}. {FIGURE_CAPTIONS[number]}")
    run.bold = True


def add_markdown_content(document, text):
    in_equation = False
    equation_lines = []
    equation_number = 0
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if line == r"\[":
            in_equation = True
            equation_lines = []
            continue
        if line == r"\]" and in_equation:
            equation_number += 1
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            append_omml(paragraph, " ".join(equation_lines))
            paragraph.add_run(f"    ({equation_number})")
            in_equation = False
            continue
        if in_equation:
            equation_lines.append(line)
            continue
        figure_match = re.fullmatch(r"\[\[FIGURE:(\d+)\]\]", line)
        if figure_match:
            add_figure(document, int(figure_match.group(1)))
            continue
        table_match = re.fullmatch(r"\[\[TABLE:(\d+)\]\]", line)
        if table_match:
            add_table(document, int(table_match.group(1)))
            continue
        if not line:
            continue
        if line.startswith("# "):
            document.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=2)
        elif line.startswith("**") and line.endswith("**"):
            paragraph = document.add_paragraph()
            paragraph.add_run(line.strip("*")).bold = True
        else:
            paragraph = document.add_paragraph()
            add_rich_text(paragraph, line)


def add_references(document, citation_order, reference_map):
    document.add_heading("References", level=1)
    sorted_keys = sorted(
        citation_order,
        key=lambda key: (
            reference_map[key]["author"].casefold(),
            reference_map[key]["year"],
            reference_map[key]["reference"].casefold(),
        ),
    )
    for key in sorted_keys:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.add_run(reference_map[key]["reference"])


def write_tables_docx():
    document = Document()
    configure_document(document)
    document.add_heading("Editable tables", level=0)
    for number in TABLE_FILES:
        add_table(document, number)
        if number != max(TABLE_FILES):
            document.add_page_break()
    path = SUBMISSION / "Tables_editable.docx"
    document.save(path)
    return path


def add_ppt_title(slide, text):
    box = slide.shapes.add_textbox(PptInches(0.55), PptInches(0.2), PptInches(12.2), PptInches(0.55))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = PptPt(24)
    paragraph.font.bold = True
    paragraph.font.color.rgb = PptRGBColor(23, 50, 77)


def write_figures_overview():
    presentation = Presentation()
    presentation.slide_width = PptInches(13.333)
    presentation.slide_height = PptInches(7.5)
    blank = presentation.slide_layouts[6]
    for number in FIGURE_FILES:
        slide = presentation.slides.add_slide(blank)
        add_ppt_title(slide, f"Figure {number}")
        image = ROOT / "figures" / FIGURE_FILES[number]
        slide.shapes.add_picture(
            str(image),
            PptInches(1.2),
            PptInches(0.95),
            width=PptInches(10.9),
            height=PptInches(5.35),
        )
        box = slide.shapes.add_textbox(
            PptInches(0.8), PptInches(6.4), PptInches(11.8), PptInches(0.8)
        )
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = FIGURE_CAPTIONS[number]
        paragraph.font.size = PptPt(11)
        paragraph.alignment = PP_ALIGN.CENTER
    path = SUBMISSION / "Figures_overview.pptx"
    presentation.save(path)
    return path


def create_graphical_abstract(values):
    width, height = 1600, 640
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    bold_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
    title_font = ImageFont.truetype(bold_path, 48)
    box_font = ImageFont.truetype(bold_path, 31)
    text_font = ImageFont.truetype(font_path, 24)
    draw.text((800, 42), "Redesign before substitution", font=title_font, fill="#17324d", anchor="mm")
    boxes = [
        (60, 145, 480, 490, "#1f5a7a", "ILLUSTRATIVE WIDTH", f"{format_value('num', values['baseline_width_mm'])}→{format_value('num', values['primary_candidate_width_mm'])} mm ceiling:\n{format_value('pct', values['primary_width_zero_compensation_saving'])}\nBreak-even length ratio:\n{format_value('num', values['primary_width_break_even_longitudinal_ratio'])}×"),
        (590, 145, 1010, 490, "#2f7d57", "SELECTABLE WIDTH", f"Half of tasks suitable:\n{format_value('pct', values['center_mid_task_mix_saving'])}\nbefore penalties"),
        (1120, 145, 1540, 490, "#d17a22", "SHORTER STRAW", f"{format_value('pct', values['straw_primary_length_reduction'])} shorter break-even:\n{format_value('pct', values['straw_primary_expected_replacement_rate_break_even'])} expected additional\nreplacement rate\n{format_value('pct', values['straw_primary_failure_probability_break_even'])} independent failure\nprobability"),
    ]
    for x1, y1, x2, y2, color, heading, body in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=28, fill=color)
        draw.text(((x1 + x2) / 2, y1 + 58), heading, font=box_font, fill="white", anchor="mm")
        draw.multiline_text(
            ((x1 + x2) / 2, y1 + 190),
            body,
            font=text_font,
            fill="white",
            anchor="mm",
            align="center",
            spacing=12,
        )
    draw.text(
        (800, 565),
        "Benefit remains conditional on function, response, manufacturing, and logistics.",
        font=text_font,
        fill="#4d5b66",
        anchor="mm",
    )
    png = SUBMISSION / "Graphical_abstract.png"
    image.save(png, dpi=(300, 300))
    presentation = Presentation()
    presentation.slide_width = PptInches(13.333)
    presentation.slide_height = PptInches(5.333)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(
        PptInches(0.5), PptInches(0.25), PptInches(12.333), PptInches(0.6)
    )
    title_paragraph = title.text_frame.paragraphs[0]
    title_paragraph.text = "Redesign before substitution"
    title_paragraph.font.size = PptPt(30)
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = PptRGBColor(23, 50, 77)
    title_paragraph.alignment = PP_ALIGN.CENTER
    for index, (_, _, _, _, color, heading, body) in enumerate(boxes):
        left = 0.5 + index * 4.42
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            PptInches(left),
            PptInches(1.2),
            PptInches(3.5),
            PptInches(2.85),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = PptRGBColor.from_string(color.lstrip("#"))
        shape.line.fill.background()
        text_frame = shape.text_frame
        text_frame.clear()
        text_frame.margin_left = PptInches(0.15)
        text_frame.margin_right = PptInches(0.15)
        heading_paragraph = text_frame.paragraphs[0]
        heading_paragraph.text = heading
        heading_paragraph.font.size = PptPt(18)
        heading_paragraph.font.bold = True
        heading_paragraph.font.color.rgb = PptRGBColor(255, 255, 255)
        heading_paragraph.alignment = PP_ALIGN.CENTER
        body_paragraph = text_frame.add_paragraph()
        body_paragraph.text = body
        body_paragraph.font.size = PptPt(15)
        body_paragraph.font.color.rgb = PptRGBColor(255, 255, 255)
        body_paragraph.alignment = PP_ALIGN.CENTER
        body_paragraph.space_before = PptPt(14)
    footer = slide.shapes.add_textbox(
        PptInches(0.6), PptInches(4.55), PptInches(12.133), PptInches(0.4)
    )
    footer_paragraph = footer.text_frame.paragraphs[0]
    footer_paragraph.text = (
        "Benefit remains conditional on function, response, manufacturing, "
        "and logistics."
    )
    footer_paragraph.font.size = PptPt(14)
    footer_paragraph.font.color.rgb = PptRGBColor(77, 91, 102)
    footer_paragraph.alignment = PP_ALIGN.CENTER
    pptx = SUBMISSION / "Graphical_abstract_editable.pptx"
    presentation.save(pptx)
    return png, pptx


def write_highlights():
    highlights = [
        "Geometry reduction is tested per successful service, not per item.",
        "Toilet-paper width benefit fails at elasticity minus one.",
        "Center perforation can match width to heterogeneous task demand.",
        "Straw shortening is optimized jointly with container geometry.",
        "Packing and rebound create conditional, discontinuous benefits.",
    ]
    text_path = SUBMISSION / "Highlights.txt"
    text_path.write_text("\n".join(f"• {item}" for item in highlights) + "\n", encoding="utf-8")
    document = Document()
    configure_document(document)
    document.add_heading("Highlights", level=0)
    for item in highlights:
        document.add_paragraph(item, style="List Bullet")
    docx_path = SUBMISSION / "Highlights.docx"
    document.save(docx_path)
    return highlights, text_path, docx_path


def write_title_page():
    document = Document()
    configure_document(document)
    add_title_page(document)
    path = SUBMISSION / "Title_page.docx"
    document.save(path)
    return path


def write_cover_letter():
    document = Document()
    configure_document(document)
    document.add_heading("Cover letter draft", level=0)
    paragraphs = [
        "[AUTHOR ACTION BEFORE SUBMISSION: insert submission date.]",
        "Editors-in-Chief\nJournal of Cleaner Production",
        "Dear Editors-in-Chief,",
        (
            "Please consider our Original article, “Redesign before substitution: "
            "functional geometry optimization of disposable products,” for the "
            "Journal of Cleaner Production."
        ),
        (
            "The manuscript presents a reproducible multi-case quantitative "
            "modeling and scenario-analysis framework for preventing material "
            "use through product geometry. Toilet-paper width, selectable center "
            "perforation, and straw–container co-design are evaluated using exact "
            "break-even conditions, transparent uncertainty, and integer logistics. "
            "The analysis emphasizes conditional and falsifiable benefit rather "
            "than assuming that nominal lightweighting produces environmental gain."
        ),
        (
            "The study fits the journal’s prevention-oriented cleaner-production "
            "scope by connecting upstream product redesign to material efficiency, "
            "consumer response, manufacturing yield, and packaging. It is neither "
            "a human-use trial nor a full life-cycle assessment; these boundaries "
            "are explicit throughout."
        ),
        (
            "Code, processed data, acquisition ledgers, tests, figures, tables, "
            "and value traceability are supplied in a reproducible repository. "
            "The manuscript makes no patentability or freedom-to-operate claim."
        ),
        (
            "[AUTHOR ACTION BEFORE SUBMISSION: confirm that this manuscript is "
            "original, has not been published, is not under consideration elsewhere, "
            "and all authors approve its submission.]"
        ),
        (
            "Sincerely,\nTatsuki Onishi\n"
            "[AUTHOR ACTION: affiliation]\n"
            "[AUTHOR ACTION: corresponding-author email]"
        ),
    ]
    for text in paragraphs:
        document.add_paragraph(text)
    path = SUBMISSION / "Cover_letter.docx"
    document.save(path)
    return path


def write_supplement():
    document = Document()
    configure_document(document)
    document.add_heading("Supplementary material", level=0)
    sections = [
        (
            "S1. Evidence classification",
            "Each parameter is labeled OBSERVED, SOURCED, CALCULATED, or "
            "ASSUMED_SCENARIO. The source, assumption, acquisition, and data "
            "dictionaries distributed with the repository provide field-level "
            "provenance. Scenario outputs do not become empirical evidence through "
            "simulation.",
        ),
        (
            "S2. Analytical break-even identities",
            "For width reduction with unchanged basis mass and process yield, "
            "material ratio is the product of width and successful-event length "
            "ratios. The exact length break-even is the inverse width ratio and "
            "the constant-elasticity break-even is minus one. For straw shortening "
            "with zero shifted burden, the expected additional replacement-rate "
            "break-even is the reduction divided by retained length. Under an "
            "independent per-attempt failure model, expected attempts are one "
            "divided by success probability and the zero-shift failure boundary "
            "equals the length-reduction fraction. Shifted container burden is "
            "defined per successful serving and added once. Center-perforation "
            "benefit is calculated from a task mixture and expected half-units, "
            "not assumed to be 50%.",
        ),
        (
            "S3. Scenario distributions",
            (
                ROOT / "config/scenarios.json"
            ).read_text(encoding="utf-8"),
        ),
        (
            "S4. Stress tests",
            "The pipeline tests zero and strong compensation, zero logistics "
            "benefit, unchanged packaging, integer-packing inefficiency, process "
            "penalties, low acceptance, high replacement, alternative functional "
            "units, conservative functional minima, and removal of PFAS evidence. "
            "The one-way sensitivity table and full fixed-seed draws are provided "
            "as machine-readable files.",
        ),
        (
            "S5. Proposed toilet-paper validation",
            "Use a randomized balanced crossover of baseline and candidate geometry. "
            "Measure roll mass or instrumented dispensing at the event level. "
            "Prespecify successful-event material as primary outcome and record "
            "length, units, folding, repeat use, adequacy, failure, and acceptability. "
            "Cross width and center perforation in a 2×2 design and measure selection "
            "and cross-tear errors. Ethics review is required before human enrollment.",
        ),
        (
            "S6. Proposed straw validation",
            "Use a geometric bench rig varying container height, fill level, "
            "insertion point, angle, beverage viscosity, and straw length. Record "
            "accessible liquid, residual mass, protrusion, handling force, buckling, "
            "leakage, detachment, and replacement. Only candidates meeting bench "
            "criteria should proceed to randomized user evaluation.",
        ),
        (
            "S7. Reproduction",
            "Create a Python environment from requirements.txt and run `make all`. "
            "The build regenerates calculated data, figures, tables, manuscript, "
            "submission files, and quality-control reports. Fixed random seeds are "
            "stored in config/scenarios.json.",
        ),
    ]
    for heading, body in sections:
        document.add_heading(heading, level=1)
        for block in body.split("\n"):
            if block.strip():
                document.add_paragraph(block)
    path = SUBMISSION / "Supplementary_material.docx"
    document.save(path)
    return path


def convert_to_pdf(docx_path):
    subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(docx_path.parent),
            str(docx_path),
        ],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    return docx_path.with_suffix(".pdf")


def write_manifest(paths):
    manifest = SUBMISSION / "submission_manifest.csv"
    with manifest.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["file", "purpose"])
        for path, purpose in paths:
            writer.writerow([path.name, purpose])
    return manifest


def write_zip():
    archive = ROOT / "functional_geometry_JCLP_submission_FINAL.zip"
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for path in sorted(SUBMISSION.iterdir()):
            if path.is_file():
                bundle.write(path, arcname=path.name)
        for suffix in ("*.png", "*.pdf", "*.svg"):
            for path in sorted((ROOT / "figures").glob(suffix)):
                bundle.write(path, arcname=f"separate_figures/{path.name}")
    return archive


def main():
    SUBMISSION.mkdir(parents=True, exist_ok=True)
    obsolete_figure_deck = SUBMISSION / "Figures_editable.pptx"
    if obsolete_figure_deck.exists():
        obsolete_figure_deck.unlink()
    values = load_values()
    references = json.loads(REFERENCES.read_text(encoding="utf-8"))
    text = substitute_values(TEMPLATE.read_text(encoding="utf-8"), values)
    text, citation_order = resolve_citations(text, references)
    document = Document()
    configure_document(document)
    add_title_page(document)
    add_markdown_content(document, text)
    add_references(document, citation_order, references)
    manuscript = SUBMISSION / "Manuscript.docx"
    document.save(manuscript)
    manuscript_pdf = convert_to_pdf(manuscript)
    tables = write_tables_docx()
    figures = write_figures_overview()
    graphical_png, graphical_pptx = create_graphical_abstract(values)
    highlights, highlights_text, highlights_docx = write_highlights()
    title_page = write_title_page()
    cover = write_cover_letter()
    supplement = write_supplement()
    supplement_pdf = convert_to_pdf(supplement)
    manifest = write_manifest(
        [
            (manuscript, "Editable main manuscript with inline figures and tables"),
            (manuscript_pdf, "Rendered manuscript for review"),
            (title_page, "Separate editable title page"),
            (cover, "Editable cover-letter draft; confirmation fields remain"),
            (highlights_docx, "Required editable highlights"),
            (highlights_text, "Plain-text highlights"),
            (supplement, "Editable supplementary methods and protocols"),
            (supplement_pdf, "Rendered supplement for review"),
            (tables, "Separate editable tables"),
            (
                figures,
                "Overview deck containing raster previews; vector sources are separate",
            ),
            (graphical_png, "Graphical abstract raster"),
            (graphical_pptx, "Editable graphical abstract slide with native shapes"),
            *[
                (
                    ROOT / "figures" / Path(filename).with_suffix(".svg").name,
                    f"Editable vector source for Figure {number}",
                )
                for number, filename in FIGURE_FILES.items()
            ],
            *[
                (
                    ROOT / "figures" / Path(filename).with_suffix(".pdf").name,
                    f"Publication vector file for Figure {number}",
                )
                for number, filename in FIGURE_FILES.items()
            ],
            *[
                (
                    ROOT / "figures" / filename,
                    f"Publication raster preview for Figure {number}",
                )
                for number, filename in FIGURE_FILES.items()
            ],
        ]
    )
    archive = write_zip()
    print(
        f"Built {manuscript.name}, {len(citation_order)} references, "
        f"{len(highlights)} highlights, and {archive.name}."
    )


if __name__ == "__main__":
    main()
